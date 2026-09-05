"""Exportacion a Excel: una hoja por seccion del README principal §16.

Dos decisiones que conviene entender antes de tocar este archivo:

**Hoja 4 no trae la serie temporal completa.** El motor corre en el navegador
(ADR-001) y lo que se persiste en `simulations.metrics` son las metricas, no los
~centenares de fotogramas de concentracion por nodo. La hoja documenta lo que de
verdad esta guardado y lo dice; inventar una curva a partir de las metricas seria
reintroducir el motor en Python por la puerta de atras.

**Hoja 5 declara que no hay modelo de ML.** Es la respuesta honesta de ADR-002:
`log Kp` sale de una correlacion publicada de 1992, no de un modelo entrenado.
Una hoja titulada "Machine Learning" que dijera cualquier otra cosa seria
exactamente la credibilidad falsa que el proyecto se comprometio a no fabricar.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.report import RegulatoryCheckResponse
from app.schemas.simulation import SimulationRecord

_HEADER_FILL = PatternFill("solid", fgColor="0F172A")
_HEADER_FONT = Font(color="E2E8F0", bold=True, size=11)
_TITLE_FONT = Font(bold=True, size=13, color="0F172A")
_NOTE_FONT = Font(italic=True, size=9, color="475569")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(bottom=_THIN)

_UNAVAILABLE = "No disponible"

_LAYER_LABELS = {
    "stratum_corneum": "Estrato corneo",
    "viable_epidermis": "Epidermis viable",
    "dermis": "Dermis",
    "hypodermis": "Hipodermis",
}

_CONFIDENCE_LABELS = {"high": "Alta", "medium": "Media", "low": "Baja"}
_BAND_LABELS = {
    "low": "Bajo",
    "moderate": "Moderado",
    "high": "Alto",
    "very_high": "Muy alto",
}
_OUTCOME_LABELS = {
    "pass": "Dentro de limite",
    "attention": "Requiere atencion",
    "fail": "Fuera de limite",
    "not_applicable": "No aplica",
    "unknown": "Sin cobertura",
}


def _write_header(sheet: Worksheet, headers: list[str], row: int = 1) -> None:
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BORDER
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def _autosize(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _note(sheet: Worksheet, row: int, text: str, span: int = 6) -> int:
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = _NOTE_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    sheet.row_dimensions[row].height = 30
    return row + 1


def _sheet_formulation(book: Workbook, record: SimulationRecord, area_cm2: float) -> None:
    sheet = book.create_sheet("1. Formulacion")
    _write_header(
        sheet,
        [
            "ID",
            "Fecha",
            "Ingrediente",
            "Concentracion (% p/p)",
            "Vehiculo",
            "Zona",
            "Area (cm2)",
            "Tiempo (h)",
        ],
    )

    source = record.input_snapshot
    zone = "No especificada"
    if record.skin_model_id:
        zone = f"Modelo de piel {record.skin_model_id}"

    sheet.append(
        [
            record.id,
            record.created_at or datetime.now(UTC).isoformat(),
            source.ingredient.name,
            record.concentration_pct,
            source.vehicle.name,
            zone,
            area_cm2,
            record.duration_hours,
        ]
    )
    sheet.append([])
    sheet.append(["Titulo", record.title])
    sheet.append(["pH", record.ph])
    sheet.append(["Dosis aplicada (mg/cm2)", record.applied_dose_mg_cm2])
    sheet.append(["Version del motor", record.engine_version])
    sheet.append(["Notas del usuario", record.notes or "—"])

    _autosize(sheet, [38, 26, 26, 22, 24, 22, 14, 14])


def _sheet_properties(book: Workbook, record: SimulationRecord) -> None:
    sheet = book.create_sheet("2. Propiedades")
    _write_header(sheet, ["Ingrediente", "Propiedad", "Valor", "Unidad", "Fuente", "DOI"])

    ingredient = record.input_snapshot.ingredient
    vehicle = record.input_snapshot.vehicle
    name = ingredient.name

    rows: list[list[Any]] = [
        [name, "Peso molecular", ingredient.molecular_weight, "g/mol", "Catalogo curado", ""],
        [name, "logP", ingredient.log_p, "adimensional", "Catalogo curado", ""],
        [
            name,
            "pKa",
            ingredient.pka if ingredient.pka is not None else _UNAVAILABLE,
            "adimensional",
            "Catalogo curado",
            "",
        ],
        [name, "Clase", ingredient.category or _UNAVAILABLE, "—", "Catalogo curado", ""],
        [
            name,
            "Banderas de riesgo",
            ", ".join(ingredient.risk_flags) if ingredient.risk_flags else "ninguna",
            "—",
            "Catalogo curado",
            "",
        ],
        [
            vehicle.name,
            "Factor potenciador",
            vehicle.enhancer_factor,
            "adimensional",
            "Catalogo de vehiculos",
            "",
        ],
    ]
    for row in rows:
        sheet.append(row)

    _note(
        sheet,
        sheet.max_row + 2,
        "La procedencia detallada por campo (base, identificador, tipo y nivel del dato) vive en "
        "la columna 'sources' del catalogo de ingredientes. Un valor sin fuente declarada debe "
        "tratarse como estimado.",
    )
    _autosize(sheet, [26, 24, 18, 16, 30, 24])


def _sheet_evidence(book: Workbook, record: SimulationRecord) -> None:
    sheet = book.create_sheet("3. Evidencia")
    _write_header(
        sheet,
        [
            "Fuente",
            "Tipo de estudio",
            "Modelo experimental",
            "Concentracion",
            "Condiciones",
            "Resultado",
            "DOI",
        ],
    )

    sheet.append(
        [
            "Potts, R. O. y Guy, R. H. (1992)",
            "Correlacion QSPR sobre datos publicados",
            "Piel humana in vitro (celda de difusion)",
            "No aplica",
            "Permeabilidad piel/agua",
            "log Kp = -2.7 + 0.71 logP - 0.0061 MW",
            "10.1023/A:1015810312465",
        ]
    )
    sheet.append(
        [
            "Bos, J. D. y Meinardi, M. M. (2000)",
            "Revision",
            "Piel humana",
            "No aplica",
            "Limite de peso molecular para penetracion",
            "Regla de los 500 Da",
            "10.1034/j.1600-0625.2000.009003165.x",
        ]
    )

    _note(
        sheet,
        sheet.max_row + 2,
        "Esta hoja recoge la evidencia que sustenta el MODELO, no ensayos sobre esta formulacion "
        "concreta. La evidencia experimental especifica del activo se anade durante la curacion "
        "del catalogo.",
        span=7,
    )
    _autosize(sheet, [32, 28, 30, 18, 30, 40, 26])


def _sheet_simulation(book: Workbook, record: SimulationRecord) -> None:
    sheet = book.create_sheet("4. Simulacion")
    _write_header(
        sheet,
        [
            "Tiempo (h)",
            "Capa",
            "Concentracion (ug/cm3)",
            "Kp (cm/h)",
            "Flujo (ug/cm2/h)",
            "Penetracion (um)",
            "Fraccion permeada (%)",
        ],
    )

    metrics = record.metrics
    duration = record.duration_hours

    # Solo la epidermis viable tiene una concentracion pico persistida; las demas
    # capas se listan con su valor no disponible en lugar de un cero enganoso.
    for layer_id, label in _LAYER_LABELS.items():
        peak = metrics.peak_concentration_ve if layer_id == "viable_epidermis" else _UNAVAILABLE
        sheet.append(
            [
                duration,
                label,
                peak,
                metrics.permeability_cm_h,
                metrics.max_flux_infinite_dose,
                metrics.penetration_depth_um,
                metrics.absorbed_fraction_pct,
            ]
        )

    sheet.append([])
    sheet.append(["Metrica", "Valor", "Unidad"])
    for label, value, unit in [
        ("log Kp", metrics.log_kp, "log(cm/h)"),
        ("Permeabilidad", metrics.permeability_cm_h, "cm/h"),
        ("Flujo maximo teorico (dosis infinita)", metrics.max_flux_infinite_dose, "ug/cm2/h"),
        ("Lag time", metrics.lag_time_hours, "h"),
        ("Fraccion absorbida", metrics.absorbed_fraction_pct, "%"),
        ("Tiempo hasta el 50 %", metrics.time_to50_pct_hours, "h"),
        ("Profundidad de penetracion", metrics.penetration_depth_um, "um"),
        ("Concentracion pico en epidermis viable", metrics.peak_concentration_ve, "ug/cm3"),
        ("Indice de irritacion (HEURISTICO)", metrics.irritation_index, "0-100"),
        (
            "Banda de irritacion",
            _BAND_LABELS.get(metrics.irritation_band, metrics.irritation_band),
            "—",
        ),
        (
            "Confianza del modelo",
            _CONFIDENCE_LABELS.get(metrics.confidence, metrics.confidence),
            "—",
        ),
    ]:
        sheet.append([label, value, unit])

    if metrics.out_of_domain_reasons:
        sheet.append([])
        sheet.append(["Motivos fuera del dominio de aplicabilidad"])
        for reason in metrics.out_of_domain_reasons:
            sheet.append([reason])

    row = _note(
        sheet,
        sheet.max_row + 2,
        "El flujo maximo teorico corresponde a dosis infinita: es una cota superior comparable "
        "entre formulas, no un estado estacionario alcanzado. Con dosis finita el vehiculo se "
        "agota antes.",
        span=7,
    )
    _note(
        sheet,
        row,
        "La serie temporal completa se calcula en el navegador y no se persiste: lo guardado son "
        "las metricas derivadas. Esta hoja no reconstruye la curva a partir de ellas.",
        span=7,
    )
    _autosize(sheet, [14, 22, 24, 16, 20, 20, 24])


def _sheet_machine_learning(book: Workbook, record: SimulationRecord) -> None:
    sheet = book.create_sheet("5. Machine Learning")
    _write_header(
        sheet,
        ["Variable", "Prediccion", "Modelo", "Version", "Metrica de validacion", "Incertidumbre"],
    )

    sheet.append(
        [
            "log Kp",
            record.metrics.log_kp,
            "Correlacion Potts-Guy (1992)",
            record.engine_version,
            "No validada contra un conjunto propio",
            "No cuantificada",
        ]
    )

    _note(
        sheet,
        sheet.max_row + 2,
        "No se emplea ningun modelo de aprendizaje automatico. log Kp procede de una correlacion "
        "publicada, no de un modelo entrenado. Entrenar uno sin un conjunto de permeabilidad "
        "medido produciria una metrica de validacion enganosa (ADR-002).",
    )
    _autosize(sheet, [22, 16, 32, 18, 34, 20])


def _sheet_ai_analysis(
    book: Workbook, record: SimulationRecord, report_content: str | None
) -> None:
    sheet = book.create_sheet("6. Analisis IA")
    _write_header(
        sheet,
        ["Resultado", "Interpretacion", "Limitaciones", "Evidencia utilizada", "Riesgo predictivo"],
    )

    interpretation = (
        report_content or "No se ha generado ningun reporte de IA para esta simulacion."
    )
    limitations = "; ".join(record.metrics.out_of_domain_reasons) or (
        "Dentro del dominio de aplicabilidad declarado (MW <= 500 g/mol, logP entre -1 y 6)."
    )

    cell_row = [
        "Confianza: "
        + _CONFIDENCE_LABELS.get(record.metrics.confidence, record.metrics.confidence),
        interpretation,
        limitations,
        "Metricas del motor determinista y correlacion Potts-Guy (1992)",
        f"Indice heuristico de irritacion: {record.metrics.irritation_index}/100 "
        f"({_BAND_LABELS.get(record.metrics.irritation_band, record.metrics.irritation_band)})",
    ]
    sheet.append(cell_row)
    for column in range(1, 6):
        sheet.cell(row=2, column=column).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 260

    _note(
        sheet,
        sheet.max_row + 2,
        "El texto de esta hoja lo redacto un modelo de lenguaje a partir de los numeros del motor. "
        "El modelo no calcula valores: interpreta. El indice de irritacion es heuristico y no esta "
        "validado experimentalmente.",
        span=5,
    )
    _autosize(sheet, [24, 70, 44, 38, 32])


def _sheet_regulatory(book: Workbook, regulatory: RegulatoryCheckResponse | None) -> None:
    sheet = book.create_sheet("7. Revision regulatoria")
    _write_header(
        sheet,
        ["Jurisdiccion", "Requisito", "Resultado preliminar", "Fuente", "Observaciones"],
    )

    if regulatory is None:
        sheet.append(
            [
                "—",
                "Verificacion regulatoria",
                "No ejecutada",
                "—",
                "No se solicito verificacion regulatoria para esta exportacion.",
            ]
        )
    else:
        for finding in regulatory.findings:
            sheet.append(
                [
                    finding.jurisdiction.upper(),
                    finding.requirement,
                    _OUTCOME_LABELS.get(finding.outcome, finding.outcome),
                    finding.source,
                    finding.message,
                ]
            )
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row].height = 60

    _note(
        sheet,
        sheet.max_row + 2,
        (regulatory.disclaimer if regulatory else "")
        or "Resultado preliminar basado en reglas con fuentes citadas. No sustituye la evaluacion "
        "de seguridad de una persona cualificada.",
        span=5,
    )
    _autosize(sheet, [16, 42, 22, 46, 70])


def _sheet_cover(book: Workbook, record: SimulationRecord) -> None:
    sheet = book.active
    sheet.title = "0. Portada"

    sheet["A1"] = "DERMASENSE — Reporte de simulacion"
    sheet["A1"].font = _TITLE_FONT

    rows = [
        ("Simulacion", record.title),
        ("Identificador", record.id),
        ("Generado", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")),
        ("Version del motor", record.engine_version),
        ("Activo", record.input_snapshot.ingredient.name),
        ("Vehiculo", record.input_snapshot.vehicle.name),
    ]
    for index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=index, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=index, column=2, value=value)

    _note(
        sheet,
        len(rows) + 5,
        "Este documento es un insumo tecnico para revision profesional. El sistema estima bajo "
        "supuestos declarados; no valida, no garantiza ni asegura la seguridad de una formulacion.",
    )
    _autosize(sheet, [26, 60])


def build_workbook(
    record: SimulationRecord,
    *,
    report_content: str | None = None,
    regulatory: RegulatoryCheckResponse | None = None,
    area_cm2: float = 10.0,
) -> bytes:
    """Genera el libro de 7 hojas (mas portada) y lo devuelve en memoria."""
    book = Workbook()

    _sheet_cover(book, record)
    _sheet_formulation(book, record, area_cm2)
    _sheet_properties(book, record)
    _sheet_evidence(book, record)
    _sheet_simulation(book, record)
    _sheet_machine_learning(book, record)
    _sheet_ai_analysis(book, record, report_content)
    _sheet_regulatory(book, regulatory)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
